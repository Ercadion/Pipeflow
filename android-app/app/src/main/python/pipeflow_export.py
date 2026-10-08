"""
측정 데이터 → 엑셀(.xlsx) — 기업 담당자가 바로 열어 필터·피벗하기 좋게 정리
앱(측정 데이터 → 엑셀 내보내기, zip 내보내기에도 포함)과 개발자 PC 도구(tools/export_excel.py)가 같은 코드를 씀

시트 (모든 시트는 '측정ID' 로 연결)
  ① 측정요약     1행 = 측정 1건 (고쳤으면 고친 결과 기준)
  ② 자동vs수정   1행 = 고친 측정 1건 × 항목 1개
  ③ 유속분포     1행 = 측정선 1개 (폭 방향 표면유속)
  ④ 촬영조건     1행 = 측정 1건 (기기·해상도·카메라 기울기·검출 방식·무결성)
  ⑤ 입력이력     1행 = 장소·실측값·메모 입력 1회
  ⑥ 데이터사전   1행 = 열 1개 (설명·단위·영문 키·추가된 양식 버전)
  ⑦ 파일정보     양식 버전·생성 시각·측정 건수·기간

확장 규칙 (회사 쪽 피벗·수식·연결이 깨지지 않게)
  - 열 이름·순서는 바꾸지 않음. 새 항목은 각 시트 '맨 뒤'에 열 추가 + COLUMNS 의 added 에 양식 버전 기록
  - 없어진 항목도 열은 남기고 빈 값
  - 양식을 바꾸면 SCHEMA_VERSION 올림
"""
from __future__ import annotations

import csv
import datetime as _dt
import json
import os
import re

from pipeflow_records import read_session, waterline

SCHEMA_VERSION = "1.0"

# (영문 키, 열 이름, 단위, 숫자 서식, 설명, 추가된 양식 버전)
SUMMARY = [
    ("id", "측정ID", "", "@", "측정 이름 = 촬영날짜_촬영시각_v앱버전. 모든 시트의 연결 키", "1.0"),
    ("date", "날짜", "", "yyyy-mm-dd", "촬영 날짜", "1.0"),
    ("time", "시각", "", "hh:mm:ss", "촬영 시각 (폰 시간)", "1.0"),
    ("lat", "위도", "°", "0.000000", "촬영 위치 위도 (GPS, 못 잡았으면 빈칸)", "1.0"),
    ("lon", "경도", "°", "0.000000", "촬영 위치 경도 (GPS, 못 잡았으면 빈칸)", "1.0"),
    ("place", "장소명", "", "@", "앱에서 저장할 때 입력한 측정 장소 (나중에 고쳤으면 최신 값)", "1.0"),
    ("pipe_id", "관로ID", "", "@", "관로 식별 번호 (예정 — 지금은 빈칸)", "1.0"),
    ("app_version", "앱 버전", "", "@", "촬영한 앱 버전", "1.0"),
    ("status", "상태", "", "@", "저장 / 저장 전(앱에서 아직 확인 중) / 예전 형식", "1.0"),
    ("corrected", "수정 여부", "", "@", "예 = 저장 전 확인 단계에서 사람이 수면선·테두리를 고침 (값은 고친 결과)", "1.0"),
    ("diameter_mm", "내경", "mm", "0", "관 내경 (설정값)", "1.0"),
    ("depth_mm", "수심", "mm", "0.0", "수심 (관 바닥 ~ 수면)", "1.0"),
    ("fill_pct", "충만율", "%", "0.0", "수심 ÷ 내경", "1.0"),
    ("v_surface", "표면유속", "m/s", "0.000", "영상(H-STIV)으로 잰 수면 유속", "1.0"),
    ("v_mean", "평균유속", "m/s", "0.000", "표면유속 × 보정계수 (단면 평균)", "1.0"),
    ("q_video", "유량(영상)", "L/s", "0.000", "평균유속 × 물이 찬 단면적 (부피 유량)", "1.0"),
    ("q_formula", "유량(공식)", "L/s", "0.000", "수심 + 경사·재질로 계산한 등류 공식 유량 (경사 입력 시)", "1.0"),
    ("truth_depth", "실측 수심", "mm", "0.0", "자 등으로 직접 잰 수심 (입력한 경우)", "1.0"),
    ("truth_v", "실측 평균유속", "m/s", "0.000", "유속계 등으로 잰 평균유속 (입력한 경우)", "1.0"),
    ("truth_q", "실측 유량", "L/s", "0.000", "유량계 등으로 잰 유량 (입력한 경우)", "1.0"),
    ("err_depth", "수심 오차", "mm", "+0.0;-0.0;0.0", "수심 − 실측 수심", "1.0"),
    ("err_v", "평균유속 오차", "m/s", "+0.000;-0.000;0.000", "평균유속 − 실측 평균유속", "1.0"),
    ("err_q", "유량 오차", "L/s", "+0.000;-0.000;0.000", "유량(영상) − 실측 유량", "1.0"),
    ("conf", "수면선 신뢰도", "", "0.00", "자동 수면선 검출 신뢰도 (0~1, 낮으면 확인 필요)", "1.0"),
    ("n_warn", "경고 수", "", "0", "분석 경고 개수 (내용은 촬영조건 시트)", "1.0"),
    ("note", "메모", "", "@", "앱에서 입력한 메모", "1.0"),
]

COMPARE = [
    ("id", "측정ID", "", "@", "", "1.0"),
    ("item", "항목", "", "@", "비교 항목", "1.0"),
    ("unit", "단위", "", "@", "", "1.0"),
    ("auto", "자동", "", "0.000", "사람이 고치기 전 자동 결과", "1.0"),
    ("corr", "수정", "", "0.000", "사람이 고친 결과", "1.0"),
    ("diff", "차이", "", "+0.000;-0.000;0.000", "수정 − 자동", "1.0"),
    ("how", "고친 방법", "", "@", "수면선 직접 맞춤 / 테두리 재검출", "1.0"),
]

PROFILE = [
    ("id", "측정ID", "", "@", "", "1.0"),
    ("line", "선 번호", "", "0", "폭 방향 측정선 (왼쪽부터 1)", "1.0"),
    ("l_mm", "폭 방향 위치", "mm", "0.0", "관 중심선 기준 가로 위치 (− 왼쪽, + 오른쪽)", "1.0"),
    ("v_mps", "표면유속", "m/s", "0.000", "그 선의 표면유속", "1.0"),
    ("method", "방법", "", "@", "C = 상관, F = 주파수 (H-STIV 내부 방식)", "1.0"),
]

CAPTURE = [
    ("id", "측정ID", "", "@", "", "1.0"),
    ("device", "기기", "", "@", "촬영 폰", "1.0"),
    ("width", "해상도 가로", "px", "0", "분석에 쓴 사진 크기", "1.0"),
    ("height", "해상도 세로", "px", "0", "", "1.0"),
    ("frames", "프레임 수", "", "0", "유속 분석용 영상 프레임 수", "1.0"),
    ("fps", "fps", "", "0.0", "실제 촬영 속도", "1.0"),
    ("f_px", "초점거리", "px", "0", "카메라 초점거리 (없으면 약투시 근사)", "1.0"),
    ("zoom", "줌", "×", "0.0", "촬영 배율", "1.0"),
    ("tilt", "카메라 기울기", "°", "0.0", "관 끝단 면과 카메라 시선 사이 각", "1.0"),
    ("roll", "수면선 기울기", "°", "0.0", "영상 속 수면선 기울기", "1.0"),
    ("method", "계산 방식", "", "@", "완전투시 / 약투시", "1.0"),
    ("rim_source", "테두리 검출", "", "@", "촬영 화면 포착 / 전체 영상 자동 / 고친 테두리", "1.0"),
    ("inner_source", "내경 테두리", "", "@", "끝단 면 안쪽 / 테두리 하나 등", "1.0"),
    ("slope_pct", "경사", "%", "0.00", "설정한 관 경사 (등류 공식용)", "1.0"),
    ("material", "재질", "", "@", "설정한 관 재질", "1.0"),
    ("warnings", "경고", "", "@", "분석 경고 내용", "1.0"),
    ("integrity", "무결성", "", "@", "저장 후 파일이 바뀌지 않았는지 (정상 / 불일치:<파일> / 기록 없음)", "1.0"),
    ("folder", "폴더", "", "@", "zip·앱 저장소 안의 측정 폴더", "1.0"),
    ("app_build", "앱 커밋", "", "@", "앱을 빌드한 커밋", "1.0"),
]

HISTORY = [
    ("id", "측정ID", "", "@", "", "1.0"),
    ("when", "입력 시각", "", "yyyy-mm-dd hh:mm:ss", "값을 입력·수정한 시각 (예전 형식은 빈칸)", "1.0"),
    ("field", "항목", "", "@", "장소명 / 실측 수심 / 실측 평균유속 / 실측 유량 / 메모", "1.0"),
    ("value", "값", "", "General", "입력한 값 (빈칸 = 지움)", "1.0"),
    ("app_version", "앱 버전", "", "@", "입력할 때의 앱 버전", "1.0"),
]

SHEETS = [("측정요약", SUMMARY), ("자동vs수정", COMPARE), ("유속분포", PROFILE),
          ("촬영조건", CAPTURE), ("입력이력", HISTORY)]

FIELD_NAMES = dict(place="장소명", truth_depth_mm="실측 수심[mm]", truth_v_mean_mps="실측 평균유속[m/s]",
                   truth_Q_Lps="실측 유량[L/s]", note="메모")


def _width(v):
    """엑셀 열 너비용 글자 폭 (한글 등은 2칸, 날짜·시각은 고정)"""
    if v is None:
        return 0
    if isinstance(v, (_dt.date, _dt.time)):
        return 10
    if isinstance(v, float):
        return len(f"{v:.6g}")
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(v))


def _header(c):
    return f"{c[1]}[{c[2]}]" if c[2] else c[1]


def _num(x):
    try:
        v = float(x)
        return v if v == v and abs(v) != float("inf") else None
    except (TypeError, ValueError):
        return None


def _sub(a, b):
    return None if a is None or b is None else a - b


def session_id(path):
    """앱과 같은 측정 이름: 날짜 폴더 안이면 '날짜_시각_v버전', 아니면 폴더 이름"""
    path = path.rstrip("/\\")
    parent = os.path.basename(os.path.dirname(path))
    name = os.path.basename(path)
    if re.fullmatch(r"\d{8}", parent) and not os.path.exists(os.path.join(os.path.dirname(path), "meta.json")):
        return f"{parent}_{name}", f"{parent}/{name}"
    return name, name


def _when(sid):
    m = re.match(r"(\d{4})(\d{2})(\d{2})_(\d{2})(\d{2})(\d{2})", sid)
    if not m:
        return None, None
    y, mo, d, h, mi, s = map(int, m.groups())
    try:
        return _dt.date(y, mo, d), _dt.time(h, mi, s)
    except ValueError:
        return None, None


def _rows_for(path):
    sid, folder = session_id(path)
    S = read_session(path)
    meta, auto, corr, ann = S["meta"], S["auto"], S["corrected"], S["ann"]
    fin = corr if corr is not None else (auto or {})
    ok = bool(fin.get("ok"))
    lv = (fin.get("level") or {}) if ok else {}
    st = (fin.get("velocity_stiv") or {}) if ok else {}
    vf = (fin.get("velocity_formula") or {}) if ok else {}
    alv = ((auto or {}).get("level") or {})
    loc = meta.get("location") or {}
    params = meta.get("params") or {}
    date, time = _when(sid)
    warns = []
    for src in (lv.get("warnings"), st.get("warnings"), vf.get("warnings"), fin.get("warnings")):
        warns += [str(w) for w in (src or [])]
    if not ok and fin:
        warns.append("분석 실패: " + str(fin.get("error", "")))
    depth, v_mean, q = _num(lv.get("depth_mm")), _num(st.get("v_mean_mps")), _num(st.get("Q_Lps"))
    t_depth, t_v, t_q = _num(ann.get("truth_depth_mm")), _num(ann.get("truth_v_mean_mps")), _num(ann.get("truth_Q_Lps"))
    fill = _num(lv.get("fill_ratio"))
    status = "저장" if S["record"] is not None else ("저장 전" if S["pending"] else "예전 형식")
    summary = dict(
        id=sid, date=date, time=time, lat=_num(loc.get("lat")), lon=_num(loc.get("lon")),
        place=ann.get("place"), pipe_id=None,
        app_version=meta.get("app_version") or ((re.search(r"_v([\w.\-]+)$", sid) or [None, ""])[1]).split("_")[0],
        status=status, corrected="예" if corr is not None else "아니오",
        diameter_mm=_num(lv.get("diameter_mm")) or _num(params.get("diameter_mm")),
        depth_mm=depth, fill_pct=None if fill is None else fill * 100,
        v_surface=_num(st.get("v_surface_mps")), v_mean=v_mean, q_video=q, q_formula=_num(vf.get("Q_Lps")),
        truth_depth=t_depth, truth_v=t_v, truth_q=t_q,
        err_depth=_sub(depth, t_depth), err_v=_sub(v_mean, t_v), err_q=_sub(q, t_q),
        conf=_num(alv.get("waterline_confidence")), n_warn=len(warns), note=ann.get("note"))

    compare = []
    if corr is not None and auto is not None and auto.get("ok") and corr.get("ok"):
        k = S["correction"] or {}
        how = ", ".join(x for x, on in (("수면선 직접 맞춤", k.get("waterline_pts")), ("테두리 재검출", k.get("rim_redetect"))) if on) \
            or ("예전 형식" if S["format"] != "v0.3.6" else "")
        A = (auto.get("level") or {}, auto.get("velocity_stiv") or {}, auto.get("velocity_formula") or {})
        C = (corr.get("level") or {}, corr.get("velocity_stiv") or {}, corr.get("velocity_formula") or {})
        items = [("수심", "mm", 0, "depth_mm", 1), ("충만율", "%", 0, "fill_ratio", 100), ("수면선 기울기", "°", 0, "roll_deg", 1),
                 ("카메라 기울기", "°", 0, "tilt_deg", 1), ("내경", "mm", 0, "diameter_mm", 1),
                 ("표면유속", "m/s", 1, "v_surface_mps", 1), ("평균유속", "m/s", 1, "v_mean_mps", 1),
                 ("유량(영상)", "L/s", 1, "Q_Lps", 1), ("유량(공식)", "L/s", 2, "Q_Lps", 1)]
        for name, unit, part, key, sc in items:
            a, c = _num(A[part].get(key)), _num(C[part].get(key))
            if a is None and c is None:
                continue
            a = None if a is None else a * sc
            c = None if c is None else c * sc
            compare.append(dict(id=sid, item=name, unit=unit, auto=a, corr=c, diff=_sub(c, a), how=how))
        wa, wc = waterline(auto), waterline(corr)
        for i, nm in enumerate(("수면선 왼쪽 x", "수면선 왼쪽 y", "수면선 오른쪽 x", "수면선 오른쪽 y")):
            if wa and wc:
                compare.append(dict(id=sid, item=nm, unit="px", auto=wa[i], corr=wc[i], diff=wc[i] - wa[i], how=how))

    profile = [dict(id=sid, line=i + 1, l_mm=_num(l.get("l_mm")), v_mps=_num(l.get("v_mps")), method=l.get("used"))
               for i, l in enumerate(sorted(st.get("lines") or [], key=lambda l: l.get("l_mm", 0)))]

    stl = meta.get("still") or {}
    fr = meta.get("frames") or {}
    integ = S["integrity"]
    capture = dict(
        id=sid, device=meta.get("device"), width=stl.get("width"), height=stl.get("height"),
        frames=fr.get("count"), fps=_num(st.get("fps")) or _num(fr.get("fps")),
        f_px=_num((meta.get("intrinsics") or {}).get("f_px")), zoom=_num(meta.get("zoom_ratio")),
        tilt=_num(lv.get("tilt_deg")), roll=_num(lv.get("roll_deg")),
        method={"perspective": "완전투시", "weak_perspective": "약투시"}.get(lv.get("method"), lv.get("method")),
        rim_source={"live_seed": "촬영 화면 포착", "auto_ransac": "전체 영상 자동",
                    "manual_or_previous": "고친 테두리"}.get(lv.get("rim_source"), lv.get("rim_source")),
        inner_source={"inner_ring": "끝단 면 안쪽", "detected_is_inner": "안쪽 테두리", "single_rim": "테두리 하나"}
        .get(lv.get("inner_source"), lv.get("inner_source")),
        slope_pct=None if _num(params.get("slope")) is None else _num(params.get("slope")) * 100,
        material=params.get("material"), warnings=" / ".join(warns) or None,
        integrity="정상" if integ == "ok" else ("기록 없음" if integ == "no_record" else integ.replace("mismatch:", "불일치:")),
        folder=folder, app_build=meta.get("app_build"))

    history = []
    for e in S["ann_log"]:
        t = e.get("time_ms")
        history.append(dict(id=sid, when=_local(t).replace(microsecond=0) if t else None,
                            field=FIELD_NAMES.get(e.get("field"), e.get("field")), value=e.get("value"),
                            app_version=e.get("app_version")))
    return summary, compare, profile, capture, history


def _decimals(fmt):
    """숫자 서식의 소수 자리 + 1 (화면은 서식대로, 값은 한 자리 더 보존)"""
    m = re.match(r"[+]?0\.(0+)", fmt)
    return len(m.group(1)) + 1 if m else (1 if fmt.lstrip("+").startswith("0") else None)


def _round_rows(rows, cols):
    for r in rows:
        for c in cols:
            v, d = r.get(c[0]), _decimals(c[3])
            if isinstance(v, float) and d is not None:
                r[c[0]] = round(v, d)


_TZ = None   # 앱에서 폰의 시간대를 넘겨 줌 (Python 이 안드로이드 시간대를 모를 수 있음)


def _local(ms):
    if _TZ is None:
        return _dt.datetime.fromtimestamp(ms / 1000)
    return (_dt.datetime.fromtimestamp(ms / 1000, _dt.timezone.utc) + _TZ).replace(tzinfo=None)


def collect(paths):
    """측정 폴더 경로들 → 시트별 행 목록 (측정ID 순)"""
    data = {name: [] for name, _ in SHEETS}
    errors = []
    for p in paths:
        try:
            s, c, pr, ca, h = _rows_for(p)
        except Exception as e:      # 손상된 측정 하나 때문에 전체가 실패하지 않게
            errors.append(f"{os.path.basename(p)}: {type(e).__name__}: {e}")
            continue
        data["측정요약"].append(s); data["자동vs수정"] += c; data["유속분포"] += pr
        data["촬영조건"].append(ca); data["입력이력"] += h
    for name, cols in SHEETS:
        data[name].sort(key=lambda r: str(r.get("id")))
        _round_rows(data[name], cols)
    return data, errors


def write_xlsx(paths, out_path, created_by="", csv_dir=None):
    """측정 폴더들 → 엑셀 파일. csv_dir 를 주면 시트별 CSV(UTF-8 BOM)도. 반환: dict(n, errors)"""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    data, errors = collect(paths)
    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True)
    for ti, (name, cols) in enumerate(SHEETS):
        ws = wb.create_sheet(name)
        ws.append([_header(c) for c in cols])
        for r in data[name]:
            ws.append([r.get(c[0]) for c in cols])
        n = len(data[name])
        for j, c in enumerate(cols, 1):
            L = get_column_letter(j)
            ws.cell(1, j).font = head_font
            ws.cell(1, j).alignment = Alignment(wrap_text=True, vertical="center")
            for i in range(2, n + 2):
                ws.cell(i, j).number_format = c[3]
            w = max([_width(_header(c)) + 2] + [_width(ws.cell(i, j).value) + 2 for i in range(2, min(n, 300) + 2)])
            ws.column_dimensions[L].width = min(max(w, 8), 45)
        ws.freeze_panes = "B2"
        if n:
            ref = f"A1:{get_column_letter(len(cols))}{n + 1}"
            t = Table(displayName=f"T{ti + 1}", ref=ref)
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        # 눈에 띄게: 고친 측정, 낮은 신뢰도, 무결성 이상
        keys = [c[0] for c in cols]
        if name == "측정요약" and n:
            last = n + 1
            yellow = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
            red = PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")
            cc = get_column_letter(keys.index("corrected") + 1)
            ws.conditional_formatting.add(f"{cc}2:{cc}{last}", CellIsRule(operator="equal", formula=['"예"'], fill=yellow))
            cf = get_column_letter(keys.index("conf") + 1)
            ws.conditional_formatting.add(f"{cf}2:{cf}{last}", FormulaRule(formula=[f'AND(ISNUMBER({cf}2),{cf}2<0.5)'], fill=red))
        if name == "촬영조건" and n:
            ci = get_column_letter(keys.index("integrity") + 1)
            ws.conditional_formatting.add(f"{ci}2:{ci}{n + 1}", FormulaRule(
                formula=[f'LEFT({ci}2,3)="불일치"'], fill=PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")))

    ws = wb.create_sheet("데이터사전")
    ws.append(["시트", "열", "단위", "영문 키", "설명", "추가된 양식 버전"])
    for name, cols in SHEETS:
        for c in cols:
            ws.append([name, _header(c), c[2], c[0], c[4], c[5]])
    for j, w in enumerate((12, 22, 8, 14, 70, 14), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
        ws.cell(1, j).font = head_font
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("파일정보")
    dates = [r["date"] for r in data["측정요약"] if r.get("date")]
    now_ms = _dt.datetime.now(_dt.timezone.utc).timestamp() * 1000
    info = [("양식 버전", SCHEMA_VERSION), ("만든 시각", _local(now_ms).replace(microsecond=0)),
            ("만든 곳", created_by), ("측정 건수", len(data["측정요약"])),
            ("기간 시작", min(dates) if dates else None), ("기간 끝", max(dates) if dates else None),
            ("고친 측정", sum(1 for r in data["측정요약"] if r["corrected"] == "예")),
            ("실측값 있는 측정", sum(1 for r in data["측정요약"] if r.get("truth_depth") is not None
                                    or r.get("truth_v") is not None or r.get("truth_q") is not None)),
            ("읽지 못한 측정", "; ".join(errors) or "없음"),
            ("안내", "모든 시트는 '측정ID' 로 연결됩니다. 열 설명·단위는 '데이터사전' 시트")]
    for k, v in info:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60
    ws["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["B5"].number_format = ws["B6"].number_format = "yyyy-mm-dd"
    for i in range(1, len(info) + 1):
        ws.cell(i, 1).font = head_font
    wb.active = 0
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

    if csv_dir:
        os.makedirs(csv_dir, exist_ok=True)
        for name, cols in SHEETS:
            with open(os.path.join(csv_dir, f"{name}.csv"), "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow([_header(c) for c in cols])
                for r in data[name]:
                    w.writerow(["" if r.get(c[0]) is None else r.get(c[0]) for c in cols])
    return dict(n=len(data["측정요약"]), errors=errors)


def export_json(paths_json, out_path, created_by="", csv_dir="", tz_offset_min=None):
    """Kotlin(Chaquopy) 에서 호출: 측정 폴더 경로 JSON 배열 → 엑셀. tz_offset_min: 폰 시간대(분). 반환: 결과 JSON 문자열"""
    global _TZ
    try:
        if tz_offset_min is not None:
            _TZ = _dt.timedelta(minutes=int(tz_offset_min))
        res = write_xlsx(json.loads(paths_json), out_path, created_by, csv_dir or None)
        return json.dumps(dict(ok=True, **res), ensure_ascii=False)
    except Exception as e:
        import traceback
        return json.dumps(dict(ok=False, error=f"{type(e).__name__}: {e}", traceback=traceback.format_exc()),
                          ensure_ascii=False)
